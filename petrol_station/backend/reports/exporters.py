"""Render a report dict (see builder.py) to PDF or Excel bytes."""
import io
from datetime import date, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND = "#0f766e"


def fmt(value, kind, currency=""):
    if value is None or value == "":
        return "-"
    if kind == "money":
        return f"{currency} {Decimal(str(value)):,.2f}".strip()
    if kind == "litres":
        return f"{Decimal(str(value)):,.2f} L"
    if kind == "number":
        return f"{Decimal(str(value)):,.2f}".rstrip("0").rstrip(".") if isinstance(value, (Decimal, float)) else str(value)
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    return str(value)


FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_text(text):
    """Stop spreadsheet formula injection: text typed by users (e.g. an expense description
    '=HYPERLINK(...)') must never run as a formula when the Excel file is opened."""
    text = str(text)
    return "'" + text if text.startswith(FORMULA_START) else text


def _cell_value(value, kind):
    if value is None:
        return None
    if kind in ("money", "litres", "number") and not isinstance(value, str):
        return float(value)
    if isinstance(value, (date, datetime)):
        return value
    return safe_text(value)


def to_excel(report, station_name, currency):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BRAND[1:])
    ws["A1"] = safe_text(f"{station_name} — {report['title']}")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = safe_text(report.get("subtitle", ""))
    ws["A3"] = f"Generated {datetime.now():%d/%m/%Y %H:%M}"
    row = 5
    for label, value, kind in report["summary"]:
        ws.cell(row=row, column=1, value=safe_text(label)).font = Font(bold=True)
        cell = ws.cell(row=row, column=2, value=_cell_value(value, kind))
        if kind == "money":
            cell.number_format = f'"{currency}" #,##0.00'
        elif kind == "litres":
            cell.number_format = '#,##0.00 "L"'
        row += 1
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 24

    for index, table in enumerate(report["tables"], start=1):
        sheet = wb.create_sheet(title=f"{index}. {table['title']}"[:31].replace("/", "-"))
        for col, (_, label, _) in enumerate(table["columns"], start=1):
            c = sheet.cell(row=1, column=col, value=label)
            c.font, c.fill, c.alignment = bold, head_fill, Alignment(horizontal="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(col)].width = max(12, min(40, len(label) + 6))
        for r, data in enumerate(table["rows"], start=2):
            for col, (key, _, kind) in enumerate(table["columns"], start=1):
                c = sheet.cell(row=r, column=col, value=_cell_value(data.get(key), kind))
                if kind == "money":
                    c.number_format = "#,##0.00"
                elif kind == "litres":
                    c.number_format = "#,##0.000"
                elif kind == "date":
                    c.number_format = "dd/mm/yyyy"
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(report, station_name, currency):
    buffer = io.BytesIO()
    wide = any(len(t["columns"]) > 6 for t in report["tables"])
    page = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(buffer, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=report["title"])
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7, leading=8.5)
    story = [
        Paragraph(f"<b>{escape(station_name)}</b>", styles["Title"]),
        Paragraph(escape(report["title"]), styles["Heading2"]),
        Paragraph(escape(report.get("subtitle", "")), styles["Normal"]),
        Paragraph(f"Generated {datetime.now():%d/%m/%Y %H:%M}", small),
        Spacer(1, 6 * mm),
    ]
    if report["summary"]:
        data = [[label, fmt(value, kind, currency)] for label, value, kind in report["summary"]]
        t = Table(data, colWidths=[70 * mm, 60 * mm], hAlign="LEFT")
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story += [t, Spacer(1, 6 * mm)]

    usable = page[0] - 24 * mm
    for table in report["tables"]:
        story.append(Paragraph(table["title"], styles["Heading3"]))
        if not table["rows"]:
            story += [Paragraph("No records for this period.", styles["Normal"]), Spacer(1, 4 * mm)]
            continue
        header = [Paragraph(f"<b>{label}</b>", small) for _, label, _ in table["columns"]]
        body = [[Paragraph(escape(fmt(row.get(key), kind, "" if kind == "money" else currency)), small)
                 for key, _, kind in table["columns"]] for row in table["rows"]]
        t = Table([header] + body, repeatRows=1, colWidths=[usable / len(header)] * len(header))
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story += [t, Spacer(1, 6 * mm)]
    doc.build(story)
    return buffer.getvalue()
