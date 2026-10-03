"""Security checklist: each test proves one item of the list is in place."""
import io

import pytest
from openpyxl import load_workbook
from werkzeug.security import check_password_hash

from backend.app import create_app
from backend.config import TestConfig
from backend.extensions import db
from backend.models import User
from backend.reports.exporters import safe_text, to_excel

from .conftest import browser_client, login

PUBLIC_API = {("POST", "/api/auth/login"), ("POST", "/api/auth/logout"), ("GET", "/api/demo/info"),
              ("POST", "/api/demo/reset"), ("POST", "/api/demo/login")}   # demo ones: 403 off the training portal


def _url(rule):
    url = rule.rule
    for arg in rule.arguments:
        url = (url.replace(f"<int:{arg}>", "1").replace(f"<{arg}>", "daily")
               .replace(f"<any(activate, deactivate):{arg}>", "activate")
               .replace(f"<any(approve, reject, void, cancel):{arg}>", "approve")
               .replace(f"<any(disable, enable):{arg}>", "disable")
               .replace(f"<any(remove, restore):{arg}>", "remove"))
    return url


# 6. Enforce server-side auth ------------------------------------------------------------
def test_every_api_route_requires_login(app):
    client = browser_client(app)
    unprotected = []
    for rule in app.url_map.iter_rules():
        if not rule.rule.startswith("/api/"):
            continue
        for method in rule.methods - {"HEAD", "OPTIONS"}:
            if (method, rule.rule) in PUBLIC_API:
                continue
            if client.open(_url(rule), method=method, json={}).status_code != 401:
                unprotected.append((method, rule.rule))
    assert unprotected == []


def test_roles_are_enforced_on_the_server(app, manager):
    manager.post("/api/users", json={"full_name": "Asha", "username": "asha01", "password": "Temp1234"})
    asha = login(app, "asha01", "Temp1234")
    for method, url in (("GET", "/api/users"), ("GET", "/api/stock/current"), ("GET", "/api/capital"),
                        ("GET", "/api/audit"), ("POST", "/api/prices"), ("GET", "/api/director/overview"),
                        ("GET", "/api/money/day"), ("PUT", "/api/settings")):
        assert asha.open(url, method=method, json={}).status_code == 403, url
    assert manager.get("/api/director/overview").status_code == 403


# 7. Lock record access (no reading other stations' / other people's records by id) --------
def test_records_of_another_station_cannot_be_opened_by_id(app, manager):
    director = login(app, "director", "Director123!")
    other = director.post("/api/director/stations", json={"name": "Other Station"}).json["station"]
    director.post("/api/director/managers", json={"full_name": "Otto", "username": "otto", "password": "Temp1234",
                                                  "station_ids": [other["id"]]})
    otto = login(app, "otto", "Temp1234")
    otto_shift = otto.get("/api/shifts").json["current_shift_id"]
    otto_pump = otto.get("/api/pumps").json["pumps"][0]["id"]
    # The first station's manager guesses the other station's ids:
    assert manager.get(f"/api/shifts/{otto_shift}").status_code == 404
    assert manager.post(f"/api/shifts/{otto_shift}/start").status_code == 404
    assert manager.put(f"/api/pumps/{otto_pump}", json={"pump_name": "x"}).status_code == 404
    assert manager.post(f"/api/shifts/{otto_shift}/deposit", json={"actual_amount": 1}).status_code == 404


def test_staff_only_see_their_own_refills_and_expenses(app, manager):
    for u in ("asha01", "baraka02"):
        manager.post("/api/users", json={"full_name": u, "username": u, "password": "Temp1234"})
    asha, baraka = login(app, "asha01", "Temp1234"), login(app, "baraka02", "Temp1234")
    day = asha.get("/api/time").json["now"][:10]
    e = asha.post("/api/expenses", json={"category_name": "Water", "description": "x", "amount": 100,
                                         "expense_date": day}).json["expense"]
    assert baraka.get("/api/expenses").json["expenses"] == []
    assert baraka.put(f"/api/expenses/{e['id']}", json={"category_name": "Water", "description": "y",
                                                        "amount": 1}).status_code == 403


# 8. Block field tampering ------------------------------------------------------------------
def test_extra_fields_cannot_change_roles_or_status(app, manager):
    r = manager.post("/api/users", json={"full_name": "Sly", "username": "sly01", "password": "Temp1234",
                                         "role": "DIRECTOR", "is_active": True})
    assert r.status_code == 403
    manager.post("/api/users", json={"full_name": "Sly", "username": "sly01", "password": "Temp1234",
                                     "must_change_password": False, "password_hash": "x", "station_id": 1})
    sly = login(app, "sly01", "Temp1234")
    me = sly.get("/api/auth/me").json["user"]
    assert me["role"] == "STAFF" and me["must_change_password"] is True
    day = sly.get("/api/time").json["now"][:10]
    e = sly.post("/api/expenses", json={"category_name": "Water", "description": "x", "amount": 100,
                                        "expense_date": day, "status": "APPROVED", "approved_by": 1}).json["expense"]
    assert e["status"] == "PENDING" and e["approved_by"] is None
    assert sly.put("/api/settings", json={"settings": {"currency": "USD"}}).status_code == 403


# 9. Secure session cookies -----------------------------------------------------------------
def test_session_cookie_flags(app):
    r = browser_client(app).post("/api/auth/login", json={"username": "manager", "password": "ChangeMe123!"})
    cookie = r.headers["Set-Cookie"]
    assert "HttpOnly" in cookie and "SameSite=Lax" in cookie


def test_forged_cross_site_requests_are_blocked(app):
    plain = app.test_client()            # like a form on another website: no JSON, no header
    login(app, "manager", "ChangeMe123!")
    r = plain.post("/api/auth/login", data={"username": "manager", "password": "ChangeMe123!"})
    assert r.status_code == 403 and "did not come from this website" in r.json["error"]


# 10. Hash passwords --------------------------------------------------------------------------
def test_passwords_are_hashed(app):
    with app.app_context():
        user = User.query.filter_by(username="manager").first()
        assert user.password_hash != "ChangeMe123!" and user.password_hash.startswith(("scrypt:", "pbkdf2:"))
        assert check_password_hash(user.password_hash, "ChangeMe123!")
    r = login(app, "manager", "ChangeMe123!").get("/api/auth/me").json["user"]
    assert "password_hash" not in r


# 11. Rate limit login --------------------------------------------------------------------------
def test_login_is_rate_limited_per_device_and_account(app):
    client = browser_client(app)
    codes = [client.post("/api/auth/login", json={"username": "manager", "password": "wrong"}).status_code
             for _ in range(6)]
    assert codes[:5] == [401] * 5 and codes[5] == 429          # account locked after 5 wrong passwords
    many = [client.post("/api/auth/login", json={"username": f"nobody{i}", "password": "x"}).status_code
            for i in range(20)]
    assert 429 in many                                          # device limit across many usernames


# 12. Bot protection ------------------------------------------------------------------------------
def test_bots_filling_the_hidden_field_are_refused(app):
    r = browser_client(app).post("/api/auth/login", json={"username": "manager", "password": "ChangeMe123!",
                                                          "website": "http://spam.example"})
    assert r.status_code == 401


# 14. Validate all input -----------------------------------------------------------------------------
def test_invalid_values_get_a_clear_400_not_a_crash(app, manager):
    assert manager.post("/api/prices", json={"fuel_type_id": "abc", "price_per_litre": 10}).status_code == 400
    assert manager.post("/api/prices", json={"fuel_type_id": 1, "price_per_litre": -5}).status_code == 400
    assert manager.get("/api/audit?limit=lots").status_code == 400
    assert manager.post("/api/users", json={"full_name": "x", "username": "<script>", "password": "Temp1234"}).status_code == 400


# 15. Escape user content (spreadsheet formulas) --------------------------------------------------------
def test_excel_export_never_runs_user_text_as_a_formula():
    assert safe_text("=HYPERLINK(\"http://evil\")").startswith("'=")
    report = {"title": "T", "subtitle": "", "summary": [], "tables": [{"title": "Expenses",
              "columns": [("d", "Description", "text")], "rows": [{"d": "=1+1"}, {"d": "@SUM(A1)"}, {"d": "Soap"}]}]}
    ws = load_workbook(io.BytesIO(to_excel(report, "=Evil Station", "TZS"))).worksheets[1]
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == ["'=1+1", "'@SUM(A1)", "Soap"]


# 16. Restrict file uploads --------------------------------------------------------------------------------
def _upload(client, name, content):
    return client.post("/api/attachments", data={"file": (io.BytesIO(content), name)},
                       content_type="multipart/form-data")


def test_only_real_pdfs_and_photos_can_be_uploaded(app, manager):
    assert _upload(manager, "evil.html", b"<script>alert(1)</script>").status_code == 400
    assert _upload(manager, "evil.png", b"<html><script>alert(1)</script></html>").status_code == 400  # fake photo
    assert _upload(manager, "tricky.pdf", b"\x89PNG\r\n\x1a\n0000").status_code == 400                # wrong type
    ok = _upload(manager, "receipt.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64)
    assert ok.status_code == 201
    file = manager.get(f"/api/attachments/{ok.json['attachment']['id']}")
    assert file.mimetype == "image/png" and "sandbox" in file.headers["Content-Security-Policy"]
    assert "stored_name" not in ok.json["attachment"]


# 17. Trim API responses --------------------------------------------------------------------------------------
def test_responses_do_not_leak_secrets_or_internals(app, manager):
    users = manager.get("/api/users").json["users"]
    for u in users:
        assert not {"password_hash", "failed_login_count", "locked_until"} & set(u)
    manager.post("/api/users", json={"full_name": "Asha", "username": "asha01", "password": "Temp1234"})
    asha = login(app, "asha01", "Temp1234")
    assert all(set(t) == {"id", "name", "fuel_type_id", "fuel"} for t in asha.get("/api/tanks").json["tanks"])
    r = asha.get("/api/this-does-not-exist")
    assert r.status_code == 404 and "Traceback" not in r.get_data(as_text=True)


# 18. Security headers + 19. Force HTTPS ----------------------------------------------------------------
def test_security_headers(app):
    h = app.test_client().get("/login.html").headers
    for name in ("Content-Security-Policy", "X-Content-Type-Options", "X-Frame-Options", "Referrer-Policy",
                 "Permissions-Policy", "Cross-Origin-Opener-Policy"):
        assert name in h, name
    assert "frame-ancestors 'self'" in h["Content-Security-Policy"] and "object-src 'none'" in h["Content-Security-Policy"]


class HttpsConfig(TestConfig):
    FORCE_HTTPS = True
    SESSION_COOKIE_SECURE = True


def test_https_is_forced_on_the_internet():
    app = create_app(HttpsConfig)
    try:
        client = app.test_client()
        r = client.get("/login.html", base_url="http://example.com")
        assert r.status_code == 308 and r.headers["Location"].startswith("https://example.com/")
        r = client.get("/login.html", base_url="https://example.com")
        assert r.status_code == 200 and "max-age=31536000" in r.headers["Strict-Transport-Security"]
        login_r = client.post("/api/auth/login", json={"username": "manager", "password": "ChangeMe123!"},
                              base_url="https://example.com")
        assert "Secure" in login_r.headers["Set-Cookie"]
    finally:
        with app.app_context():
            db.drop_all()


# 1. Hide API keys: production refuses to run without its own secret key ---------------------------------
class ProductionWithoutSecret(TestConfig):
    ENV_NAME = "production"
    TESTING = False
    SECRET_KEY_FROM_ENV = False


def test_production_needs_a_secret_key_from_the_environment():
    with pytest.raises(RuntimeError, match="SECRET_KEY"):
        create_app(ProductionWithoutSecret)
